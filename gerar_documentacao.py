# -*- coding: utf-8 -*-
"""Gera documento Word com todas as melhorias do projeto VetClinic."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

OUTPUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "Documentacao_Melhorias_VetClinic.docx",
)


def add_title(doc, text):
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_h1(doc, text):
    doc.add_heading(text, level=1)


def add_h2(doc, text):
    doc.add_heading(text, level=2)


def add_h3(doc, text):
    doc.add_heading(text, level=3)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * level)
    for run in p.runs:
        run.font.size = Pt(11)


def add_numbered(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        run.font.size = Pt(11)


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(0)


def build_document():
    doc = Document()

    # Margens
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    add_title(doc, "VetClinic — Documentação Completa das Melhorias")
    add_para(
        doc,
        "Sistema de Gestão para Clínica Veterinária | Trabalho Acadêmico Flutter + SQLite",
        bold=True,
    )
    doc.add_paragraph()
    add_para(doc, "Documento gerado para registrar todas as alterações, melhorias e implementações realizadas no projeto original baixado via Git, desde o estado inicial até a versão funcional atual.")
    doc.add_paragraph()

    # --- 1 ---
    add_h1(doc, "1. Objetivo deste documento")
    add_para(
        doc,
        "Este arquivo descreve, de forma detalhada, tudo o que foi modificado no projeto "
        "original da Clínica Veterinária. O objetivo é permitir que você (e a banca avaliadora) "
        "tenham dimensão completa do trabalho realizado: o que existia antes, o que foi criado, "
        "o que foi corrigido e como o sistema funciona hoje."
    )

    # --- 2 ---
    add_h1(doc, "2. Como era o projeto ORIGINAL (antes das alterações)")
    add_para(doc, "Ao abrir o projeto baixado do Git, a situação era a seguinte:", bold=True)

    add_h2(doc, "2.1. O que já existia (pontos positivos)")
    add_bullet(doc, "Estrutura modular em pastas: lib/modules/ com separação model, repository, controller e page.")
    add_bullet(doc, "11 módulos CRUD criados: Cliente, Animal, Profissional, Consulta, Prontuário, Medicamento, Vacina, Prescrição, Aplicação de Vacina, Fatura e Item de Fatura.")
    add_bullet(doc, "Schema SQLite (schema.dart) com tabelas relacionadas e chaves estrangeiras definidas.")
    add_bullet(doc, "DatabaseHelper configurado com sqflite e sqflite_common_ffi (necessário para Windows).")
    add_bullet(doc, "Operações básicas de insert, update, delete e listagem nos repositories.")
    add_bullet(doc, "pubspec.yaml com dependências sqflite, sqflite_common_ffi e path.")

    add_h2(doc, "2.2. Problemas e lacunas do projeto original")
    add_bullet(doc, "O main.dart apenas exibia a mensagem \"SQLite OK\" — o app não tinha interface utilizável.")
    add_bullet(doc, "Nenhuma navegação entre telas: os CRUDs existiam no código, mas não eram acessíveis pelo usuário.")
    add_bullet(doc, "Não havia tela de login nem controle de acesso.")
    add_bullet(doc, "Não havia menu lateral ou organização visual das funcionalidades.")
    add_bullet(doc, "As listagens mostravam apenas IDs (ex.: id_cliente) em vez de nomes e informações úteis.")
    add_bullet(doc, "Formulários exigiam digitar IDs manualmente para relacionamentos (ex.: id_cliente ao cadastrar animal).")
    add_bullet(doc, "Datas e valores digitados como texto livre, sem validação nem seletores.")
    add_bullet(doc, "Não existia pesquisa por ID centralizada.")
    add_bullet(doc, "Não havia dashboard, relatórios, alertas nem regras de negócio do enunciado.")
    add_bullet(doc, "SQLite não havia sido testado de ponta a ponta na interface.")
    add_bullet(doc, "Sem dados de demonstração (seed) para apresentação.")
    add_bullet(doc, "Sem README ou documentação técnica.")
    add_bullet(doc, "Sem testes automatizados relevantes (apenas teste padrão do Flutter Demo).")

    # --- 3 ---
    add_h1(doc, "3. Primeira fase de melhorias (estrutura base do sistema)")
    add_para(doc, "Na primeira intervenção, foi construída a fundação do aplicativo funcional:", bold=True)

    add_h2(doc, "3.1. Banco de dados reorganizado (versão 2)")
    add_bullet(doc, "Schema expandido para 14 tabelas, incluindo: usuario, sala_atendimento e pagamento.")
    add_bullet(doc, "Campos adicionais em tabelas existentes: preferências de contato (cliente), peso/histórico/vacinação (animal), sala e lembrete em consultas, estoque mínimo e alerta de interação (medicamento), observações em fatura.")
    add_bullet(doc, "Foreign keys ativadas com PRAGMA foreign_keys = ON.")
    add_bullet(doc, "Migração automática onUpgrade para bancos antigos.")
    add_bullet(doc, "Usuário admin inicial e salas padrão inseridos no seed.")

    add_h2(doc, "3.2. Sistema de navegação e layout")
    add_bullet(doc, "Criação do MainShell com menu lateral fixo (280px), agrupado em: Principal, Cadastros, Atendimento, Farmácia e Financeiro.")
    add_bullet(doc, "AppRoutes centralizando todas as rotas e títulos das páginas.")
    add_bullet(doc, "Tema visual VetClinic (cores verde veterinário, cards, inputs estilizados) em app_theme.dart.")
    add_bullet(doc, "Barra superior com título da seção atual e indicador \"Sistema online\".")

    add_h2(doc, "3.3. Autenticação")
    add_bullet(doc, "Tela de LoginPage com layout dividido (banner + formulário).")
    add_bullet(doc, "AuthService validando e-mail e senha contra tabela usuario.")
    add_bullet(doc, "Credenciais demo: admin@clinica.com / admin123.")

    add_h2(doc, "3.4. Telas principais criadas")
    add_bullet(doc, "DashboardPage — contadores de clientes, animais, consultas, faturas, medicamentos e vacinas.")
    add_bullet(doc, "PesquisaPage — busca por ID em qualquer tabela, com exibição de campos e registros relacionados.")
    add_bullet(doc, "SearchService — serviço unificado de pesquisa cross-table.")

    add_h2(doc, "3.5. CRUD padronizado")
    add_bullet(doc, "Widget CrudScaffold reutilizável: cabeçalho, busca, listagem em cards, formulário modal, confirmação de exclusão.")
    add_bullet(doc, "Todas as 13 páginas de módulos refatoradas para usar o CrudScaffold.")
    add_bullet(doc, "Método findById adicionado em todos os repositories.")
    add_bullet(doc, "Busca nas listagens aceita texto ou ID numérico.")

    add_h2(doc, "3.6. Novos módulos")
    add_bullet(doc, "Módulo Sala (sala_atendimento) — cadastro de salas de consultório e cirurgia.")
    add_bullet(doc, "Módulo Pagamento — registro de pagamentos vinculados a faturas.")
    add_bullet(doc, "Módulo Usuario — usado internamente para login.")

    add_h2(doc, "3.7. Correções técnicas")
    add_bullet(doc, "main.dart reescrito: inicializa FFI, abre banco e lança VetClinicApp.")
    add_bullet(doc, "Tipos corrigidos nos models: valor_total e valor como double, quantidade_estoque como int.")
    add_bullet(doc, "Nomenclatura padronizada (camelCase nos models Dart, snake_case no SQLite).")

    # --- 4 ---
    add_h1(doc, "4. Segunda fase de melhorias (polimento e requisitos do enunciado)")
    add_para(doc, "Após análise das fraquezas, foi implementado um pacote amplo de melhorias:", bold=True)

    add_h2(doc, "4.1. Experiência do usuário (UX)")
    add_bullet(doc, "Dropdowns nos formulários: cliente, animal, profissional, sala, fatura, prontuário, medicamento e vacina — sem digitar IDs.")
    add_bullet(doc, "Seletor de data (showDatePicker) para campos de data.")
    add_bullet(doc, "Seletor de data e hora (showDatePicker + showTimePicker) para consultas.")
    add_bullet(doc, "Dropdown de status e tipos: status da consulta, tipo de procedimento, status da fatura, forma de pagamento.")
    add_bullet(doc, "LookupService — cache de nomes para exibir labels amigáveis nos dropdowns.")
    add_bullet(doc, "Menu responsivo: sidebar em telas largas (≥900px), drawer (gaveta) em telas menores.")

    add_h2(doc, "4.2. Validações")
    add_bullet(doc, "validators.dart com validação de CPF (dígitos verificadores), e-mail, datas ISO e números positivos.")
    add_bullet(doc, "Validação aplicada no cadastro de clientes (CPF e e-mail).")
    add_bullet(doc, "Mensagens de erro amigáveis em SnackBar ao salvar.")

    add_h2(doc, "4.3. Segurança e autenticação")
    add_bullet(doc, "Senhas armazenadas com hash SHA-256 + salt (PasswordUtils + pacote crypto).")
    add_bullet(doc, "Compatibilidade com senhas antigas em texto puro (legado).")
    add_bullet(doc, "Provider (ChangeNotifier) substituindo AuthService estático — estado global de login.")
    add_bullet(doc, "Dois perfis de usuário: admin (acesso total) e recepcao (sem Relatório Financeiro).")
    add_bullet(doc, "Segundo usuário demo: recepcao@clinica.com / recep123.")
    add_bullet(doc, "Redirecionamento automático: logado → MainShell; deslogado → LoginPage.")

    add_h2(doc, "4.4. Novas telas funcionais")
    add_bullet(doc, "Consultas de Hoje — lista consultas do dia com nomes de animal/profissional/sala e botão \"Enviar lembrete\" (simula SMS/e-mail).")
    add_bullet(doc, "Central de Alertas — medicamentos com estoque ≤ mínimo; vacinas com reforço nos próximos 30 dias; badge no menu.")
    add_bullet(doc, "Relatório Financeiro — total faturado, recebido e pendente, com filtro por período (admin).")

    add_h2(doc, "4.5. Regras de negócio (BusinessRulesService)")
    add_bullet(doc, "Conflito de agenda: impede duas consultas no mesmo horário para o mesmo profissional ou sala.")
    add_bullet(doc, "Interação medicamentosa: bloqueia prescrição se medicamentos têm alerta cruzado no campo alerta_interacao.")
    add_bullet(doc, "Estoque insuficiente: impede prescrição se quantidade = 0.")
    add_bullet(doc, "Baixa de estoque: ao criar prescrição, reduz 1 unidade do medicamento.")
    add_bullet(doc, "Recálculo de fatura: ao salvar/excluir item_fatura, soma valores e atualiza valor_total.")
    add_bullet(doc, "Status de pagamento: ao registrar pagamento, atualiza fatura para pendente/parcial/pago conforme valor pago.")

    add_h2(doc, "4.6. Dados de demonstração (seed v3)")
    add_bullet(doc, "Banco atualizado para versão 3 com seed completo na primeira execução.")
    add_bullet(doc, "2 clientes, 3 animais, 2 profissionais, 2 consultas HOJE, prontuário, prescrição, medicamentos (um com estoque baixo), vacinas, fatura com 2 itens (R$ 195,00), aplicações de vacina com reforço próximo.")

    add_h2(doc, "4.7. Dependências adicionadas")
    add_bullet(doc, "provider ^6.1.2 — gerenciamento de estado (auth).")
    add_bullet(doc, "crypto ^3.0.6 — hash de senhas.")
    add_bullet(doc, "intl ^0.20.2 — formatação de datas e moeda (R$).")

    add_h2(doc, "4.8. Testes e documentação")
    add_bullet(doc, "test/widget_test.dart — teste de abertura na tela de login.")
    add_bullet(doc, "test/repository_test.dart — teste insert/findById de cliente no SQLite.")
    add_bullet(doc, "Testes unitários de CPF, e-mail e PasswordUtils.")
    add_bullet(doc, "README.md com arquitetura, diagrama ER (Mermaid), credenciais e instruções.")

    # --- 5 ---
    add_h1(doc, "5. Estrutura de pastas ATUAL do projeto")
    add_code_block(doc, [
        "lib/",
        "├── main.dart                    # Entrada: FFI + Provider + VetClinicApp",
        "├── core/",
        "│   ├── auth/                    # (auth_service removido → AuthProvider)",
        "│   ├── database/",
        "│   │   ├── database_helper.dart # SQLite v3, migração, seed",
        "│   │   ├── schema.dart          # 14 tabelas",
        "│   │   └── seed_data.dart       # Dados demo",
        "│   ├── navigation/app_routes.dart",
        "│   ├── providers/auth_provider.dart",
        "│   ├── services/",
        "│   │   ├── alertas_service.dart",
        "│   │   ├── business_rules_service.dart",
        "│   │   ├── lookup_service.dart",
        "│   │   └── search_service.dart",
        "│   ├── theme/app_theme.dart",
        "│   ├── utils/validators.dart, password_utils.dart",
        "│   └── widgets/crud_scaffold.dart, main_shell.dart",
        "├── modules/                     # 14 módulos CRUD",
        "│   ├── cliente/, animal/, profissional/, sala/",
        "│   ├── consulta/, prontuario/",
        "│   ├── medicamento/, vacina/, prescricao/, aplicacao_vacina/",
        "│   ├── fatura/, item_fatura/, pagamento/, usuario/",
        "└── pages/",
        "    ├── login_page.dart",
        "    ├── dashboard_page.dart",
        "    ├── pesquisa_page.dart",
        "    ├── consultas_hoje_page.dart",
        "    ├── alertas_page.dart",
        "    └── relatorio_financeiro_page.dart",
    ])

    # --- 6 ---
    add_h1(doc, "6. Banco de dados — 14 tabelas")
    tabelas = [
        ("usuario", "Login e perfis (admin/recepção)"),
        ("cliente", "Tutores/responsáveis pelos animais"),
        ("animal", "Pets vinculados a clientes"),
        ("profissional", "Veterinários e equipe"),
        ("sala_atendimento", "Salas de consultório e cirurgia"),
        ("consulta", "Agendamentos (animal + profissional + sala)"),
        ("prontuario", "Histórico médico eletrônico"),
        ("medicamento", "Estoque, lote, alerta de interação"),
        ("vacina", "Cadastro de vacinas e intervalo de reforço"),
        ("prescricao", "Medicamentos prescritos no prontuário"),
        ("aplicacao_vacina", "Vacinas aplicadas e data de reforço"),
        ("fatura", "Faturas por cliente"),
        ("item_fatura", "Itens/procedimentos de cada fatura"),
        ("pagamento", "Pagamentos recebidos por fatura"),
    ]
    for nome, desc in tabelas:
        add_bullet(doc, f"{nome} — {desc}")

    # --- 7 ---
    add_h1(doc, "7. Atendimento ao enunciado da faculdade")
    enunciado = [
        ("Gestão de Clientes e Animais", "CRUD completo + histórico médico, vacinação, preferências, dropdown de tutor."),
        ("Prontuário Eletrônico", "Módulo prontuário vinculado a animal e consulta, com prescrições e vacinas."),
        ("Agendamento de Consultas", "CRUD consultas + salas + conflito de horário + Consultas de Hoje + lembretes."),
        ("Medicamentos e Vacinas", "Estoque, alerta de interação, baixa automática, reforço vacinal na Central de Alertas."),
        ("Faturamento e Financeiro", "Faturas, itens, pagamentos, recálculo automático, relatório financeiro."),
        ("Requisitos não funcionais", "Login, perfis, hash de senha, layout organizado, responsividade, validações."),
    ]
    for req, impl in enunciado:
        add_para(doc, req, bold=True)
        add_para(doc, f"→ {impl}")
        doc.add_paragraph()

    # --- 8 ---
    add_h1(doc, "8. Menu do sistema (todas as opções)")
    menu = [
        "Principal: Dashboard | Consultas de Hoje | Central de Alertas | Relatório Financeiro* | Pesquisa por ID",
        "Cadastros: Clientes | Animais | Profissionais | Salas",
        "Atendimento: Consultas | Prontuários",
        "Farmácia: Medicamentos | Vacinas | Prescrições | Aplicação de Vacinas",
        "Financeiro: Faturas | Itens de Fatura | Pagamentos",
        "* Relatório Financeiro visível apenas para perfil admin.",
    ]
    for m in menu:
        add_bullet(doc, m)

    # --- 9 ---
    add_h1(doc, "9. Resumo quantitativo das alterações")
    add_bullet(doc, "Arquivos novos criados: aproximadamente 25+ (core, pages, seed, utils, testes, README).")
    add_bullet(doc, "Arquivos modificados: todos os módulos CRUD, main.dart, pubspec.yaml, schema, database_helper.")
    add_bullet(doc, "Tabelas SQLite: de 11 para 14.")
    add_bullet(doc, "Versão do banco: 1 → 2 → 3.")
    add_bullet(doc, "Telas acessíveis ao usuário: de 0 (só \"SQLite OK\") para 19 telas funcionais.")
    add_bullet(doc, "Pacotes Flutter adicionados: provider, crypto, intl.")

    # --- 10 ---
    add_h1(doc, "10. PASSO A PASSO — Como rodar o sistema")
    add_para(doc, "Siga os passos abaixo na ordem indicada.", bold=True)
    doc.add_paragraph()

    add_h3(doc, "Pré-requisitos")
    add_numbered(doc, "Ter o Flutter SDK instalado (versão 3.11 ou superior).")
    add_numbered(doc, "Verificar instalação abrindo o terminal e digitando: flutter doctor")
    add_numbered(doc, "Ter um dispositivo/alvo configurado: Windows desktop, emulador Android, navegador Chrome, etc.")
    doc.add_paragraph()

    add_h3(doc, "Passo 1 — Abrir o terminal na pasta do projeto")
    add_para(doc, "Navegue até a pasta raiz do projeto (onde está o arquivo pubspec.yaml):")
    add_code_block(doc, [
        "cd \"c:\\Users\\diogo\\OneDrive\\Área de Trabalho\\Trab-PV-Flutter---Clinica-Veterinaria-main\\Trab-PV-Flutter---Clinica-Veterinaria-main\"",
    ])
    add_para(doc, "(Ajuste o caminho se a pasta estiver em outro local.)")
    doc.add_paragraph()

    add_h3(doc, "Passo 2 — Baixar as dependências")
    add_code_block(doc, ["flutter pub get"])
    add_para(doc, "Aguarde até concluir. Isso instala sqflite, provider, crypto, intl e demais pacotes.")
    doc.add_paragraph()

    add_h3(doc, "Passo 3 — (Recomendado) Limpar banco antigo se já rodou antes")
    add_para(
        doc,
        "Se você já executou uma versão anterior do app, o banco SQLite pode estar desatualizado. "
        "Opções: desinstalar o app e reinstalar, OU apagar o arquivo clinica_veterinaria.db "
        "(localização varia por plataforma; no Windows desktop fica na pasta de dados do app)."
    )
    add_para(doc, "Isso garante que os dados demo (seed v3) sejam criados corretamente.")
    doc.add_paragraph()

    add_h3(doc, "Passo 4 — Executar o aplicativo")
    add_para(doc, "Para Windows (recomendado para este projeto):")
    add_code_block(doc, ["flutter run -d windows"])
    add_para(doc, "Alternativas:")
    add_code_block(doc, [
        "flutter run -d chrome          # Executar no navegador",
        "flutter run -d android           # Emulador/dispositivo Android",
        "flutter devices                # Listar dispositivos disponíveis",
    ])
    doc.add_paragraph()

    add_h3(doc, "Passo 5 — Fazer login")
    add_para(doc, "Use uma das credenciais abaixo na tela de login:")
    add_code_block(doc, [
        "Administrador (acesso total):",
        "  E-mail: admin@clinica.com",
        "  Senha:  admin123",
        "",
        "Recepção (sem relatório financeiro):",
        "  E-mail: recepcao@clinica.com",
        "  Senha:  recep123",
    ])
    doc.add_paragraph()

    add_h3(doc, "Passo 6 — Explorar o sistema (roteiro sugerido para apresentação)")
    add_numbered(doc, "Dashboard — ver contadores e alertas pendentes.")
    add_numbered(doc, "Consultas de Hoje — ver consultas agendadas e testar \"Enviar lembrete\".")
    add_numbered(doc, "Central de Alertas — ver medicamento com estoque baixo (Dipirona Vet).")
    add_numbered(doc, "Clientes / Animais — cadastrar usando dropdowns (sem digitar ID).")
    add_numbered(doc, "Consultas — agendar com seletor de data/hora; tentar conflito de horário.")
    add_numbered(doc, "Prescrições — tentar prescrever Dipirona + Meloxicam (deve bloquear interação).")
    add_numbered(doc, "Itens de Fatura — adicionar item e observar total da fatura atualizar.")
    add_numbered(doc, "Pagamentos — registrar pagamento e ver status da fatura mudar.")
    add_numbered(doc, "Pesquisa por ID — buscar cliente #1 e ver animais relacionados.")
    add_numbered(doc, "Relatório Financeiro (admin) — gerar relatório por período.")
    doc.add_paragraph()

    add_h3(doc, "Passo 7 — Rodar os testes (opcional)")
    add_code_block(doc, ["flutter test"])
    add_para(doc, "Executa testes de validação, senha, login e CRUD no SQLite.")
    doc.add_paragraph()

    add_h3(doc, "Solução de problemas comuns")
    add_bullet(doc, "\"flutter não é reconhecido\" → Instale o Flutter SDK e adicione ao PATH do sistema.")
    add_bullet(doc, "Erro sqflite no Windows → O projeto já usa sqflite_common_ffi; certifique-se de rodar flutter pub get.")
    add_bullet(doc, "Tela em branco ou dados vazios → Apague o banco antigo (Passo 3) e reinicie o app.")
    add_bullet(doc, "Erro de compilação após git pull → Execute: flutter clean && flutter pub get && flutter run")
    doc.add_paragraph()

    add_h1(doc, "11. Conclusão")
    add_para(
        doc,
        "O projeto evoluiu de um protótipo com CRUDs isolados e tela \"SQLite OK\" para um sistema "
        "completo de gestão veterinária, alinhado ao enunciado acadêmico, com interface profissional, "
        "banco relacional funcional, regras de negócio, segurança básica, dados de demonstração e "
        "documentação. Todas as melhorias descritas neste documento estão implementadas no código "
        "fonte atual do repositório."
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— Fim do documento —")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(f"Documento gerado: {path}")
